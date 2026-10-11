#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Конвертер объявлений из Avito XML фида обратно в Excel (.xlsx).
Поддерживает парсинг официальных структур:
- PriceList со структурой <Service>
- Place и WorkDays со структурой <Option>
"""

import sys
import os
import argparse
import xml.etree.ElementTree as ET
import openpyxl

EXCEL_HEADERS = [
    'Id', 'ListingFee', 'AvitoId', 'ManagerName', 'ContactPhone', 'ImageUrls', 
    'ImageNames', 'ContactMethod', 'Category', 'Price', 'Title', 'Description', 
    'Address', 'ServiceType', 'ServiceSubtype', 'WorkExperience', 'Place', 
    'WorkWithContract', 'Guarantee', 'WorkDays', 'WorkTimeFrom', 'WorkTimeTo', 
    'PriceList', 'AvitoDateEnd', 'CompanyName', 'AvitoStatus', 'EMail'
]

def xml_to_excel(xml_path, xlsx_path):
    print(f"[*] Чтение XML-фида: '{xml_path}'...")
    tree = ET.parse(xml_path)
    root = tree.getroot()
    ads = root.findall('Ad')
    print(f"[+] Найдено {len(ads)} объявлений.")
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Компьютерная пом-ОС и программы'
    
    # Строка 1: Категория
    ws.cell(1, 1).value = 'Услуги - Предложение услуг - Компьютерная помощь - ОС и программы'
    
    # Строка 2: Названия параметров
    for c_idx, h in enumerate(EXCEL_HEADERS, 1):
        ws.cell(2, c_idx).value = h
        
    # Строка 3: Обязательность
    row3_map = {
        'Id': 'Обязательный', 'ListingFee': 'Необязательный', 'AvitoId': 'Необязательный',
        'ImageUrls': 'Обязательный', 'Category': 'Обязательный', 'Title': 'Обязательный',
        'Description': 'Обязательный', 'Address': 'Обязательный', 'ServiceType': 'Обязательный',
        'ServiceSubtype': 'Обязательный'
    }
    for c_idx, h in enumerate(EXCEL_HEADERS, 1):
        ws.cell(3, c_idx).value = row3_map.get(h, 'Необязательный')
        
    # Строка 4: Подсказки
    for c_idx, h in enumerate(EXCEL_HEADERS, 1):
        ws.cell(4, c_idx).value = 'Подробнее о параметре'
        
    # Заполнение данных с 5 строки
    for r_idx, ad in enumerate(ads, start=5):
        # ImageUrls
        images = []
        img_elem = ad.find('Images')
        if img_elem is not None:
            for img in img_elem.findall('Image'):
                u = img.get('url')
                if u:
                    images.append(u)
        img_urls_str = ' | '.join(images)
        
        # Place (<Option>)
        place_opts = [o.text.strip() for o in ad.findall('Place/Option') if o.text]
        place_str = ' | '.join(place_opts) if place_opts else (ad.findtext('Place') or '')
        
        # WorkDays (<Option>)
        days_opts = [o.text.strip() for o in ad.findall('WorkDays/Option') if o.text]
        days_str = ' | '.join(days_opts) if days_opts else (ad.findtext('WorkDays') or '')
        
        # PriceList (<Service>)
        pl_elem = ad.find('PriceList')
        pl_str = ''
        if pl_elem is not None:
            services = pl_elem.findall('Service')
            if services:
                lines = []
                for s in services:
                    sn = s.findtext('ServiceName') or ''
                    st = s.findtext('ServiceTitle') or ''
                    sp = s.findtext('ServicePrice') or ''
                    sf = s.findtext('ServiceStartingPrice') or 'Нет'
                    s_type = s.findtext('ServicePriceType') or 'за услугу'
                    lines.append(f"{sn}|{st}|{sp}|{sf}|{s_type}")
                pl_str = '\n'.join(lines)
            else:
                pl_str = pl_elem.text or ''
                
        for c_idx, h in enumerate(EXCEL_HEADERS, 1):
            cell = ws.cell(r_idx, c_idx)
            if h == 'ImageUrls':
                cell.value = img_urls_str
            elif h == 'Place':
                cell.value = place_str
            elif h == 'WorkDays':
                cell.value = days_str
            elif h == 'PriceList':
                cell.value = pl_str
            elif h == 'AvitoDateEnd':
                de = ad.find('DateEnd')
                cell.value = de.text if de is not None else ''
            else:
                el = ad.find(h)
                cell.value = el.text if el is not None else ''
                
    wb.save(xlsx_path)
    print(f"[+] Экспорт в Excel завершён: '{xlsx_path}' ({len(ads)} строк)")

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Конвертация XML-фида Авито в Excel (.xlsx)')
    parser.add_argument('--input', '-i', required=True, help='Входной XML-файл')
    parser.add_argument('--output', '-o', default='avito_ads_exported.xlsx', help='Выходной Excel-файл')
    args = parser.parse_args()
    
    xml_to_excel(args.input, args.output)
